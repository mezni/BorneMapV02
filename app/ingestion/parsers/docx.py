"""DOCX text parser (requires the optional ``python-docx`` dependency)."""

from __future__ import annotations

from pathlib import Path
from typing import ClassVar

from app.ingestion.parsers.base import BaseParser, ParsedDocument, ParseError

try:  # pragma: no cover - exercised only when python-docx is installed
    from docx import Document as DocxDocument  # type: ignore[import-not-found]
except ImportError:  # pragma: no cover - python-docx is optional
    DocxDocument = None


class DOCXParser(BaseParser):
    """Extracts text from Word documents (``.docx``).

    Requires ``python-docx``; install with ``uv add python-docx``.
    """

    supported_extensions: ClassVar[set[str]] = {".docx"}

    def parse(self, path: Path) -> ParsedDocument:
        if DocxDocument is None:
            raise ParseError("python-docx is not installed; install it with: uv add python-docx")
        try:
            document = DocxDocument(str(path))
        except Exception as exc:
            raise ParseError(f"Cannot read DOCX {path}: {exc}") from exc

        parts: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)

        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                parts.append(", ".join(cells))

        body = "\n".join(parts)
        core = document.core_properties
        return ParsedDocument(
            text=body,
            title=core.title or path.stem,
            metadata={
                "author": core.author,
                "created": core.created.isoformat() if core.created else None,
                "paragraph_count": len(document.paragraphs),
                "table_count": len(document.tables),
            },
        )


DOCXParser.register(DOCXParser.supported_extensions)
